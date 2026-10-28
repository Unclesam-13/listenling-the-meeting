"""手机版端到端测试：真实 Chromium + 模拟麦克风 + 模拟 Soniox / Claude 接口。

    python tests/mobile_e2e.py [站点目录，默认 mobile] [截图输出目录]

覆盖：首次设置、上传资料（md；站点里有 vendor/ 时也测 pdf 和 docx）、开始录音、
实时字幕 + 中英互译、自动识别提问并根据资料回答、改名/标记"我"、手动提问、
结束后自动生成会议总结、导出、刷新后恢复、会议记录、深色模式截图。
"""
import json
import os
import sys
import tempfile
import threading
import time
from functools import partial
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

from playwright.sync_api import sync_playwright

SITE = Path(sys.argv[1] if len(sys.argv) > 1 else Path(__file__).parent.parent / "mobile").resolve()
SHOTS = Path(sys.argv[2] if len(sys.argv) > 2 else tempfile.mkdtemp()).resolve()
SHOTS.mkdir(parents=True, exist_ok=True)
HAS_VENDOR = (SITE / "vendor" / "pdf.min.js").exists() and (SITE / "vendor" / "mammoth.browser.min.js").exists()

CORS = {"access-control-allow-origin": "*", "access-control-allow-headers": "*", "access-control-allow-methods": "POST, OPTIONS"}
LOG = {"soniox_config": None, "audio_frames": 0, "end_signal": False, "claude": []}


def serve(directory):
    class Quiet(SimpleHTTPRequestHandler):
        def log_message(self, *a):
            pass
    srv = ThreadingHTTPServer(("127.0.0.1", 0), partial(Quiet, directory=str(directory)))
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    return srv, f"http://localhost:{srv.server_address[1]}/"


def make_pdf(path, lines):
    """生成一个最简单的单页 PDF（英文，Helvetica）。"""
    content = "BT /F1 12 Tf 72 720 Td 16 TL " + " ".join(f"({l}) Tj T*" for l in lines) + " ET"
    objs = ["<< /Type /Catalog /Pages 2 0 R >>", "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
            f"<< /Length {len(content)} >>\nstream\n{content}\nendstream", "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out, offs = "%PDF-1.4\n", []
    for i, o in enumerate(objs, 1):
        offs.append(len(out.encode()))
        out += f"{i} 0 obj\n{o}\nendobj\n"
    x = len(out.encode())
    out += f"xref\n0 {len(objs)+1}\n0000000000 65535 f \n" + "".join(f"{o:010d} 00000 n \n" for o in offs)
    out += f"trailer\n<< /Size {len(objs)+1} /Root 1 0 R >>\nstartxref\n{x}\n%%EOF\n"
    Path(path).write_bytes(out.encode())


# ---------------- 模拟 Soniox ----------------
def T(text, **kw):
    return {"text": text, "is_final": True, **kw}


SCRIPT = [  # 录音开始后依次发送的识别结果
    {"tokens": [{"text": "你们的数据", "is_final": False, "speaker": "1"}]},
    {"tokens": [T("你们的数据", speaker="1", language="zh"), T("存在哪里？", speaker="1", language="zh"),
                T("Where is your data", translation_status="translation", speaker="1"),
                T(" stored?", translation_status="translation", speaker="1"), T("<end>")]},
    {"tokens": [T("We are evaluating", speaker="2", language="en"), T(" several vendors this quarter.", speaker="2", language="en"),
                T("我们本季度正在评估几家供应商。", translation_status="translation", speaker="2"), T("<end>")]},
    {"tokens": [T("How long does", speaker="2", language="en"), T(" implementation take?", speaker="2", language="en"), T("<end>")]},
    {"tokens": [T("实施需要多长时间？", translation_status="translation", speaker="2")]},  # 迟到的译文
]


def on_ws(ws):
    def handle(msg):
        if isinstance(msg, (bytes, bytearray)):
            LOG["audio_frames"] += 1
            return
        if msg == "":
            LOG["end_signal"] = True
            ws.send(json.dumps({"tokens": [T("好的，谢谢。", speaker="1", language="zh"), T("OK, thanks.", translation_status="translation", speaker="1"), T("<end>")],
                                "finished": True}))
            ws.close()
            return
        LOG["soniox_config"] = json.loads(msg)
        for m in SCRIPT:  # 在回调里直接发送（Playwright 同步接口不能跨线程调用）
            ws.send(json.dumps(m))
    ws.on_message(handle)


# ---------------- 模拟 Claude ----------------
def sse(texts):
    ev = [{"type": "message_start", "message": {"id": "m"}}]
    ev += [{"type": "content_block_delta", "index": 0, "delta": {"type": "text_delta", "text": t}} for t in texts]
    ev += [{"type": "message_stop"}]
    return "".join(f"event: {e['type']}\ndata: {json.dumps(e, ensure_ascii=False)}\n\n" for e in ev)


def on_claude(route):
    req = route.request
    if req.method == "OPTIONS":
        return route.fulfill(status=204, headers=CORS)
    body = json.loads(req.post_data or "{}")
    LOG["claude"].append(body)
    system = " ".join(b.get("text", "") for b in body.get("system", []))
    user = body["messages"][0]["content"]
    assert req.headers.get("anthropic-dangerous-direct-browser-access") == "true"
    if "判断最后一段发言" in system:  # 意图识别
        last = user.rsplit("最后一段", 1)[-1]
        q = "？" in last or "?" in last
        txt = json.dumps({"is_question": q, "confidence": 0.9, "question": last.split("：", 1)[-1].strip() if q else "", "topic": "x"}, ensure_ascii=False)
        return route.fulfill(status=200, headers={**CORS, "content-type": "application/json"},
                             body=json.dumps({"content": [{"type": "text", "text": txt}]}))
    if not body.get("stream"):
        return route.fulfill(status=200, headers={**CORS, "content-type": "application/json"},
                             body=json.dumps({"content": [{"type": "text", "text": "数据, data, AWS"}]}))
    if "会议纪要" in system:
        parts = ["## 会议摘要\n双方讨论了数据存储和实施周期。\n", "## 关键结论与决定\n- 数据存放在 AWS 悉尼区域\n",
                 "## 待办事项\n- [ ] 发送实施计划 —— 我\n", "## 对方提出的问题\n- 数据存在哪里 → AWS 悉尼\n", "## 待跟进 / 未解决\n- 无\n"]
    elif "AWS" in system or "AWS" in user:
        parts = ["**结论**：数据存放在 AWS 悉尼区域。\n", "- 传输 TLS 1.3\n", "- 静态 AES-256\n", "来源：[资料.md]"]
    else:
        parts = ["**结论**：标准实施周期 4 到 6 周。\n", "- 第 1 周需求梳理\n", "来源：[资料.md]"]
    route.fulfill(status=200, headers={**CORS, "content-type": "text/event-stream"}, body=sse(parts))


def on_deepseek(route):
    req = route.request
    if req.method == "OPTIONS":
        return route.fulfill(status=204, headers=CORS)
    body = json.loads(req.post_data or "{}")
    LOG.setdefault("deepseek", []).append(body)
    if not body.get("stream"):
        return route.fulfill(status=200, headers={**CORS, "content-type": "application/json"},
                             body=json.dumps({"choices": [{"message": {"content": "价格, price, 企业版"}}]}))
    chunks = ["**结论**：", "企业版每年 12,000 美元起。\n", "来源：[资料.md]"]
    sse_body = "".join(f"data: {json.dumps({'choices': [{'delta': {'content': c}}]}, ensure_ascii=False)}\n\n" for c in chunks) + "data: [DONE]\n\n"
    route.fulfill(status=200, headers={**CORS, "content-type": "text/event-stream"}, body=sse_body)


def on_jev(route):
    req = route.request
    if req.method == "OPTIONS":
        return route.fulfill(status=204, headers=CORS)
    body = json.loads(req.post_data or "{}")
    LOG.setdefault("jev", []).append({"body": body, "idem": req.headers.get("idempotency-key")})
    q = "?" in body["state"]["latest_utterance"]["text"] or "？" in body["state"]["latest_utterance"]["text"]
    route.fulfill(status=200, headers={**CORS, "content-type": "application/json", "x-tokens-remaining": "99000",
                                       "access-control-expose-headers": "X-Tokens-Remaining"},
                  body=json.dumps({"model": "jev-1.13.0", "answers": {"is_question": {"type": "noul", "noul": 0.95 if q else 0.05},
                                   "intent": {"type": "choice", "choice": "question" if q else "statement", "confidence": 0.9}},
                                   "usage": {"input_tokens": 250}}))


def main():
    srv, base = serve(SITE)
    errors = []
    with sync_playwright() as p:
        browser = p.chromium.launch(args=["--use-fake-ui-for-media-stream", "--use-fake-device-for-media-stream"])
        ctx = browser.new_context(viewport={"width": 390, "height": 844}, device_scale_factor=2, is_mobile=True, has_touch=True,
                                  permissions=["microphone"], locale="zh-CN")
        ctx.route_web_socket("wss://stt-rt.soniox.com/**", on_ws)
        ctx.route("https://api.soniox.com/**", lambda r: r.fulfill(status=200, headers={**CORS, "content-type": "application/json"},
                                                                    body=json.dumps({"api_key": "snx_temp_test", "expires_at": "x"})))
        ctx.route("https://api.anthropic.com/**", on_claude)
        ctx.route("https://api.deepseek.com/**", on_deepseek)
        ctx.route("https://jevmodel.org/**", on_jev)
        page = ctx.new_page()
        page.on("console", lambda m: m.type == "error" and errors.append(m.text))
        page.on("pageerror", lambda e: errors.append(str(e)))
        page.goto(base)

        # 1) 首次打开：弹出设置
        page.wait_for_selector("#settingsSheet:not([hidden])")
        page.screenshot(path=SHOTS / "01-首次设置.png")
        page.fill("input[name=sonioxKey]", "sk-soniox")
        page.fill("input[name=claudeKey]", "sk-ant-test")
        # 上传资料
        tmp = Path(tempfile.mkdtemp())
        (tmp / "资料.md").write_text("# 数据安全 Data security\n数据存储在 AWS 悉尼区域，传输 TLS 1.3，静态 AES-256。\n\n# 上线时间\n标准实施周期 4 到 6 周。\n", encoding="utf-8")
        files = [tmp / "资料.md"]
        if HAS_VENDOR:
            make_pdf(tmp / "pricing.pdf", ["Pricing", "Standard plan is 99 USD per month.", "Enterprise starts at 12000 USD per year."])
            import docx
            d = docx.Document(); d.add_heading("Roadmap 2027", 1); d.add_paragraph("Q1: Mobile app for iOS and Android."); d.save(tmp / "roadmap.docx")
            files += [tmp / "pricing.pdf", tmp / "roadmap.docx"]
        page.set_input_files("#fileInput", [str(f) for f in files])
        page.wait_for_function(f"document.querySelectorAll('#docList li').length === {len(files)}", timeout=20000)
        kb_text = page.evaluate("window.__mc.kb.docs.map(d => d.name + ':' + d.text.slice(0, 300)).join(' | ')")
        if HAS_VENDOR:
            assert "Enterprise starts" in kb_text and "Mobile app for iOS" in kb_text, kb_text
        page.screenshot(path=SHOTS / "02-设置-资料库.png", full_page=False)
        page.click("#settingsSheet [data-close]")  # 点"完成"也要保存
        page.wait_for_selector("#settingsSheet", state="hidden")
        assert page.evaluate("window.__mc.S.sonioxKey") == "sk-soniox"

        # 2) 开始录音
        page.click("#recBtn")
        page.wait_for_function("document.querySelector('#recBtn').classList.contains('on')", timeout=10000)
        page.wait_for_function("window.__mc.M.lines.length >= 3", timeout=15000)
        page.wait_for_function("window.__mc.M.answers.length >= 1 && window.__mc.M.answers.every(a => a.done)", timeout=15000)
        cfg = LOG["soniox_config"]
        assert cfg["translation"] == {"type": "two_way", "language_a": "zh", "language_b": "en"}, cfg
        assert cfg["enable_speaker_diarization"] and cfg["audio_format"] == "pcm_s16le" and "api_key" not in cfg, cfg
        assert LOG["audio_frames"] > 3, LOG["audio_frames"]
        page.screenshot(path=SHOTS / "03-回答.png")
        a0 = page.evaluate("window.__mc.M.answers[0]")
        assert "AWS 悉尼" in a0["text"] and a0["asker"] == "说话人1", a0

        # 3) 字幕 + 翻译
        page.click(".tabs button[data-tab=captions]")
        time.sleep(1.2)
        lines = page.evaluate("window.__mc.M.lines.map(l => [l.sid, l.text, l.trans])")
        assert lines[0] == ["spk1", "你们的数据存在哪里？", "Where is your data stored?"], lines
        assert lines[1][2] == "我们本季度正在评估几家供应商。", lines
        assert lines[2][2] == "实施需要多长时间？", lines  # 迟到的译文补上了
        assert page.locator("#lines .tr").count() >= 3
        page.screenshot(path=SHOTS / "04-字幕翻译.png")
        # 改名 + 标记"我"
        page.locator("#lines .who").first.click()
        page.fill("#speakerForm input[name=name]", "王总")
        page.click("#speakerForm .sheet-foot button")
        page.locator("#lines .who").nth(1).click()
        page.check("#speakerForm input[name=me]")
        page.click("#speakerForm .sheet-foot button")
        assert page.evaluate("window.__mc.M.names.spk1") == "王总" and page.evaluate("window.__mc.M.me") == ["spk2"]
        # 关闭翻译显示
        page.uncheck("#showTrans")
        assert page.locator("#tab-captions.hide-trans").count() == 1
        page.check("#showTrans")

        # 4) 手动提问
        page.fill("#askText", "实施要多久")
        page.press("#askText", "Enter")
        page.wait_for_function("window.__mc.M.answers.filter(a => a.origin === 'manual' && a.done).length === 1", timeout=10000)

        # 5) 结束 -> 自动总结
        page.click("#recBtn")
        page.wait_for_function("!document.querySelector('#recBtn').classList.contains('on')", timeout=10000)
        assert LOG["end_signal"], "没有通知 Soniox 结束"
        page.wait_for_function("window.__mc.M.summary.includes('待办事项')", timeout=15000)
        assert page.evaluate("window.__mc.M.lines.at(-1).text") == "好的，谢谢。"  # 结束时最后一句也识别完整
        page.click(".tabs button[data-tab=summary]")
        time.sleep(0.5)
        assert page.locator("#sumBody ul.todo li").count() == 1
        page.screenshot(path=SHOTS / "05-总结.png", full_page=True)
        with page.expect_download() as dl:
            page.click("#exportAll")
        exported = Path(dl.value.path()).read_text(encoding="utf-8")
        assert "## 会议全文" in exported and "（译）Where is your data stored?" in exported and "王总" in exported

        # 6) 刷新后恢复 + 会议记录
        page.reload()
        page.wait_for_function("window.__mc && window.__mc.M.lines.length >= 4", timeout=10000)
        page.click("#historyBtn")
        page.wait_for_selector("#historySheet:not([hidden])")
        assert page.locator("#histList li").count() == 1
        page.screenshot(path=SHOTS / "06-会议记录.png")
        page.click("#newMeeting")
        page.wait_for_function("window.__mc.M.lines.length === 0", timeout=5000)

        # 7) 深色模式截图（读取历史会议）
        page.emulate_media(color_scheme="dark")
        page.click("#historyBtn")
        page.locator("#histList button[data-open]").first.click()
        page.click(".tabs button[data-tab=captions]")
        time.sleep(0.3)
        page.screenshot(path=SHOTS / "07-深色-字幕.png")
        page.click(".tabs button[data-tab=answers]")
        page.screenshot(path=SHOTS / "08-深色-回答.png")

        # 8) 换成 DeepSeek 回答 + Jev 判断提问
        page.emulate_media(color_scheme="light")
        page.click("#settingsBtn")
        page.select_option("select[name=answerProvider]", "deepseek")
        page.fill("input[name=deepseekKey]", "sk-ds")
        page.select_option("#intentSel", "jev")
        page.fill("input[name=jevKey]", "jev-key")
        page.click("#settingsForm .sheet-foot button")
        page.fill("#askText", "企业版多少钱")
        page.press("#askText", "Enter")
        page.wait_for_function("window.__mc.M.answers.at(-1).done && window.__mc.M.answers.at(-1).text.includes('12,000')", timeout=10000)
        ds = LOG["deepseek"][-1]
        assert ds["model"] == "deepseek-flash" and ds["thinking"] == {"type": "disabled"} and ds["stream"], ds
        assert "Claude 未填 Key" not in page.inner_text("#statusLine")
        # 新会议，录音，让 Jev 判断
        page.click("#historyBtn"); page.click("#newMeeting")
        page.click("#recBtn")
        page.wait_for_function("window.__mc.M.answers.length >= 1 && window.__mc.M.answers.every(a => a.done)", timeout=15000)
        assert LOG["jev"] and LOG["jev"][0]["idem"] and LOG["jev"][0]["body"]["questions"]["is_question"]["type"] == "noul", LOG["jev"][:1]
        page.click("#recBtn")
        page.wait_for_function("!document.querySelector('#recBtn').classList.contains('on')", timeout=10000)

        claude_models = sorted({b["model"] for b in LOG["claude"]})
        browser.close()
    srv.shutdown()
    bad = [e for e in errors if "favicon" not in e]
    assert not bad, bad
    print(f"mobile e2e ok | 音频帧 {LOG['audio_frames']} | Claude 调用 {len(LOG['claude'])} 次 {claude_models} | "
          f"DeepSeek 调用 {len(LOG.get('deepseek', []))} 次 | Jev 调用 {len(LOG.get('jev', []))} 次 | "
          f"pdf/docx {'已测' if HAS_VENDOR else '跳过（无 vendor）'} | 截图 {SHOTS}")


if __name__ == "__main__":
    main()
