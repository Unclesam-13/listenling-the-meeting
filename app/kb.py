"""资料库：读取文件 -> 切块 -> 检索。不需要下载任何模型。

两种模式（config.yaml 的 kb.mode）：
  full    资料不多时，把全部资料直接交给 Claude（配合提示词缓存，最准确）
  search  资料较多时，用关键词检索（BM25）挑出相关片段；Claude 会先把问题扩展成中英文关键词，
          所以中文问题也能搜到英文资料
  auto    默认：资料总字数不超过 kb.full_max_chars 时用 full，否则用 search
支持 .md .txt .pdf .docx。
"""
from __future__ import annotations

import logging
import re
from dataclasses import dataclass
from pathlib import Path

import numpy as np

log = logging.getLogger("kb")

EXTS = {".md", ".txt", ".pdf", ".docx"}


@dataclass
class Chunk:
    source: str
    text: str
    page: int | None = None


def read_file(p: Path) -> list[tuple[str, int | None]]:
    """返回 [(文本, 页码)]"""
    ext = p.suffix.lower()
    if ext in (".md", ".txt"):
        return [(p.read_text(encoding="utf-8", errors="ignore"), None)]
    if ext == ".pdf":
        from pypdf import PdfReader
        return [(pg.extract_text() or "", i + 1) for i, pg in enumerate(PdfReader(str(p)).pages)]
    if ext == ".docx":
        import docx
        d = docx.Document(str(p))
        parts = [para.text for para in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        return [("\n".join(parts), None)]
    return []


def split_text(text: str, size: int, overlap: int) -> list[str]:
    """先按标题分节（Markdown # 标题），每节内再按段落拼块；每块带上所属标题，检索更准。"""
    sections, head, body = [], "", []
    for line in text.splitlines():
        if re.match(r"^#{1,6}\s+\S", line):
            if "\n".join(body).strip():
                sections.append((head, "\n".join(body)))
            head, body = line.lstrip("#").strip(), []
        else:
            body.append(line)
    if "\n".join(body).strip():
        sections.append((head, "\n".join(body)))
    out = []
    for head, body in sections:
        for c in _pack(body, size, overlap):
            out.append(f"{head}\n{c}" if head else c)
    return out


def _pack(text: str, size: int, overlap: int) -> list[str]:
    paras = [x.strip() for x in re.split(r"\n\s*\n", text) if x.strip()]
    chunks, cur = [], ""
    for para in paras:
        while len(para) > size:  # 超长段落硬切
            if cur:
                chunks.append(cur); cur = ""
            chunks.append(para[:size]); para = para[size - overlap:]
        if len(cur) + len(para) + 1 <= size:
            cur = f"{cur}\n{para}" if cur else para
        else:
            if cur:
                chunks.append(cur)
            tail = cur[-overlap:] if overlap and cur else ""
            cur = f"{tail}\n{para}" if tail else para
    if cur:
        chunks.append(cur)
    return chunks


_jieba = None


def tokenize(text: str) -> list[str]:
    """中英混合分词：英文按单词，中文用 jieba。"""
    global _jieba
    text = text.lower()
    toks = re.findall(r"[a-z0-9][a-z0-9_\-\.]*", text)
    zh = "".join(re.findall(r"[一-鿿]+", text))
    if zh:
        if _jieba is None:
            import jieba
            jieba.setLogLevel(logging.WARNING)
            _jieba = jieba
        toks += [w for w in _jieba.cut_for_search(zh) if w.strip()]
    return toks


class KnowledgeBase:
    def __init__(self, cfg: dict):
        self.cfg = cfg
        self.dir = Path(cfg["dir"])
        self.chunks: list[Chunk] = []
        self.docs: list[tuple[str, str]] = []  # (文件名, 全文)
        self.bm25 = None
        self.mode = "search"

    def build(self) -> dict:
        self.dir.mkdir(parents=True, exist_ok=True)
        files = sorted(p for p in self.dir.rglob("*") if p.suffix.lower() in EXTS and not p.name.startswith("."))
        size, ov = self.cfg.get("chunk_chars", 500), self.cfg.get("chunk_overlap", 100)
        chunks: list[Chunk] = []
        docs: list[tuple[str, str]] = []
        for f in files:
            try:
                name = str(f.relative_to(self.dir))
                pages = read_file(f)
                full = "\n\n".join((f"[第{pg}页]\n" if pg else "") + t for t, pg in pages if t.strip())
                if full.strip():
                    docs.append((name, full))
                for text, page in pages:
                    for c in split_text(text, size, ov):
                        chunks.append(Chunk(name, c, page))
            except Exception as e:
                log.warning("读取失败 %s: %s", f.name, e)
        self.chunks, self.docs = chunks, docs

        from rank_bm25 import BM25Okapi
        self.bm25 = BM25Okapi([tokenize(c.text) for c in chunks]) if chunks else None

        chars = sum(len(t) for _, t in docs)
        mode = self.cfg.get("mode", "auto")
        self.mode = mode if mode in ("full", "search") else ("full" if chars <= self.cfg.get("full_max_chars", 150000) else "search")
        log.info("资料库：%d 个文件，%d 个片段，约 %d 字，模式 %s", len(files), len(chunks), chars, self.mode)
        return {"files": len(files), "chunks": len(chunks), "chars": chars, "mode": self.mode}

    def full_text(self) -> str:
        return "\n\n".join(f"<document name=\"{n}\">\n{t}\n</document>" for n, t in self.docs)

    def search(self, query: str, k: int | None = None, keywords: list[str] | None = None) -> list[dict]:
        """关键词检索。keywords 是额外的检索词（例如 Claude 给出的英文/中文同义词），会和原问题一起参与打分。"""
        k = k or self.cfg.get("top_k", 6)
        if not self.chunks or self.bm25 is None:
            return []
        q = tokenize(query) + [t for w in (keywords or []) for t in tokenize(w)]
        s = self.bm25.get_scores(q)
        top = [int(i) for i in np.argsort(-s)[:k] if s[i] > 0]
        return [{"source": self.chunks[i].source, "page": self.chunks[i].page,
                 "text": self.chunks[i].text, "score": round(float(s[i]), 3)} for i in top]

    def stats(self) -> dict:
        return {"chunks": len(self.chunks), "files": len(self.docs), "mode": self.mode, "dir": str(self.dir)}
